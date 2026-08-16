from pptx import Presentation
from core.document import create_document
from langsmith import traceable

@traceable(
    metadata={
        "module": "loader",
        "tool":"pptx"
    }
)
def load_pptx(file_path: str):

    presentation = Presentation(file_path)

    text = ""


    for slide in presentation.slides:

        for shape in slide.shapes:

            if hasattr(shape, "text"):
                text += shape.text + "\n"


   
    return create_document(
    
            text,
    
            {
                "source": file_path,
                "type": "pptx"
            }
    
        )